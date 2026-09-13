#!/usr/bin/env python3
"""Export the JournalCache architecture deck to PPTX for Google Slides import."""
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt
from lxml import etree
import os

HERE = os.path.dirname(os.path.abspath(__file__))
XRD = os.path.join(HERE, "slides-assets", "xrootd-logo.png")
JC = os.path.join(HERE, "slides-assets", "journalcache-logo.png")
CERN = os.path.join(HERE, "slides-assets", "cern-logo.svg.png")
OUT = os.path.join(HERE, "JournalCache-architecture-slides.pptx")

BG = RGBColor(0x07, 0x07, 0x07)
INK = RGBColor(0xF7, 0xF4, 0xEF)
MUTED = RGBColor(0xC9, 0xBF, 0xB6)
GOLD = RGBColor(0xB8, 0x92, 0x3A)
PINK = RGBColor(0xE8, 0xA0, 0xBF)
PANEL = RGBColor(0x16, 0x14, 0x16)


def set_run(run, text, size=18, color=INK, bold=False, italic=False):
    run.text = text
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.font.bold = bold
    run.font.italic = italic
    run.font.name = "Calibri"


def add_tb(slide, l, t, w, h, text, size=18, color=INK, bold=False, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    set_run(run, text, size, color, bold)
    return box


def fill_shape(shape, color):
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()


def outline_shape(shape, color, width_pt=1.25):
    shape.fill.solid()
    shape.fill.fore_color.rgb = PANEL
    shape.line.color.rgb = color
    shape.line.width = Pt(width_pt)


def bg(slide):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = BG


def logos(slide, prs):
    if os.path.isfile(XRD):
        slide.shapes.add_picture(XRD, Inches(0.35), Inches(0.18), Inches(0.55), Inches(0.55))
    if os.path.isfile(JC):
        slide.shapes.add_picture(JC, Inches(1.05), Inches(0.12), height=Inches(0.62))
    if os.path.isfile(CERN):
        slide.shapes.add_picture(CERN, Inches(12.45), Inches(0.18), Inches(0.52), Inches(0.52))
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0.82), prs.slide_width, Pt(1.5))
    fill_shape(bar, GOLD)


def footer(slide, prs, n, total):
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(7.28), prs.slide_width, Inches(0.22))
    fill_shape(bar, GOLD)
    add_tb(slide, 11.6, 7.18, 1.5, 0.28, f"{n} / {total}", size=11, color=GOLD, align=PP_ALIGN.RIGHT)


def kicker(slide, text):
    add_tb(slide, 0.55, 0.95, 12, 0.32, text.upper(), size=12, color=PINK, bold=True)


def title(slide, text, y=1.22, h=0.7, size=28):
    add_tb(slide, 0.55, y, 12.2, h, text, size=size, color=INK, bold=True)


def lead(slide, text, y=1.9, h=0.85, size=16):
    add_tb(slide, 0.55, y, 12.2, h, text, size=size, color=MUTED)


def card(slide, l, t, w, h, heading, body_lines, label=None, label_color=GOLD):
    sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(l), Inches(t), Inches(w), Inches(h))
    outline_shape(sh, GOLD, 1.0)
    y = t + 0.1
    if label:
        add_tb(slide, l + 0.12, y, w - 0.24, 0.22, label.upper(), size=10, color=label_color, bold=True)
        y += 0.22
    add_tb(slide, l + 0.12, y, w - 0.24, 0.32, heading, size=15, color=GOLD, bold=True)
    y += 0.34
    box = slide.shapes.add_textbox(Inches(l + 0.12), Inches(y), Inches(w - 0.24), Inches(h - (y - t) - 0.1))
    tf = box.text_frame
    tf.word_wrap = True
    tf.clear()
    first = True
    for line in body_lines:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.level = 0
        p.space_after = Pt(4)
        run = p.add_run()
        prefix = "• " if not line.startswith("•") else ""
        set_run(run, prefix + line, 13, INK)


def bullets(slide, l, t, w, h, lines, size=16):
    box = slide.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.clear()
    first = True
    for line in lines:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.space_after = Pt(6)
        run = p.add_run()
        set_run(run, "•  " + line, size, INK)


def table(slide, l, t, w, h, rows, col_w=None):
    n_rows, n_cols = len(rows), len(rows[0])
    shp = slide.shapes.add_table(n_rows, n_cols, Inches(l), Inches(t), Inches(w), Inches(h))
    tbl = shp.table
    if col_w:
        for i, cw in enumerate(col_w):
            tbl.columns[i].width = Inches(cw)
    for r, row in enumerate(rows):
        for c, val in enumerate(row):
            cell = tbl.cell(r, c)
            cell.text = ""
            p = cell.text_frame.paragraphs[0]
            p.alignment = PP_ALIGN.LEFT
            run = p.add_run()
            is_hdr = r == 0
            set_run(run, val, 11 if not is_hdr else 10, GOLD if is_hdr else INK, bold=is_hdr)
            cell.text_frame.word_wrap = True
            tc = cell._tc
            tcPr = tc.get_or_add_tcPr()
            solid = etree.SubElement(tcPr, qn("a:solidFill"))
            srgb = etree.SubElement(solid, qn("a:srgbClr"))
            srgb.set("val", "161416" if r else "1E1810")
    return shp


def new_slide(prs, n, total):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg(slide)
    logos(slide, prs)
    footer(slide, prs, n, total)
    return slide


def main():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    T = 16

    s = new_slide(prs, 1, T)
    if os.path.isfile(JC):
        s.shapes.add_picture(JC, Inches(0.9), Inches(1.25), height=Inches(2.07))
    add_tb(s, 3.7, 1.55, 8.6, 1.15, "JournalCache for XRootD", size=36, color=INK, bold=True)
    rule = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(3.7), Inches(2.85), Inches(1.6), Pt(3))
    fill_shape(rule, GOLD)
    add_tb(s, 3.7, 3.15, 8.6, 0.4, "Dr. Andreas-Joachim Peters - CERN IT-SD-PSS", size=20, color=PINK)
    add_tb(s, 3.7, 3.6, 8.6, 0.35, "XROOTD WORKSHOP LYON 2026", size=14, color=GOLD, bold=True)
    repo = add_tb(s, 3.7, 4.15, 8.6, 0.35, "github.com/cern-eos/xrootd - XrdClJournalCachePlugin", size=16, color=GOLD)
    repo.text_frame.paragraphs[0].runs[0].hyperlink.address = "https://github.com/cern-eos/xrootd/tree/XrdClJournalCachePlugin"

    s = new_slide(prs, 2, T)
    kicker(s, "XRootD")
    title(s, "A client read cache, distinct from XCache")
    lead(s, "JournalCache records only the byte ranges an application reads, in an append-only per-file journal. There is no block size and no read amplification beyond the request.", y=1.95, h=0.95)
    card(s, 0.55, 3.1, 3.9, 3.5, "Client or proxy", ["An XrdCl file plugin for xrdcp, ROOT, or any XrdCl application. The same plugin can run under PSS or XrdHttp. xjcd installs a forwarding proxy."], "Where", GOLD)
    card(s, 4.7, 3.1, 3.9, 3.5, "Exact ranges", ["A hit requires the full request in the journal. A miss is fetched from the origin and appended. A write resets that file’s journal."], "What", GOLD)
    card(s, 8.85, 3.1, 3.9, 3.5, "Not XCache", ["XCache (XrdPfc) is a shared, server-side block cache with prefetch. JournalCache is a client journal. The two are complementary."], "Scope", PINK)

    s = new_slide(prs, 3, T)
    kicker(s, "Talk")
    title(s, "Contents")
    bullets(s, 0.7, 2.05, 11, 4.6, [
        "1   Deployments: client and forwarding proxy",
        "2   Read path, journal format, HTTP freshness",
        "3   Origin allowlist and the xjc / xjcd / xjccleand tools",
        "4   XCache (XrdPfc) in brief",
        "5   Comparison and recommended use",
        "6   Limits and status",
    ], size=20)

    s = new_slide(prs, 4, T)
    kicker(s, "Deployments")
    title(s, "One plugin, two deployments")
    lead(s, "The file plugin is the cache. The HTTP extension and xjcd are optional: they expose the same journal as a forwarding proxy for browsers and HTTP clients.", y=1.9, h=0.75)
    card(s, 0.55, 2.8, 6.0, 3.85, "Laptop / analysis job", [
        "xrdcp · ROOT · any XrdCl application",
        "libXrdClJournalCachePlugin",
        "XRD_PLUGIN and cache=/var/tmp/journalcache/",
        "Local journal on disk; origin on miss",
    ], "Client", GOLD)
    card(s, 6.8, 2.8, 6.0, 3.85, "Forwarding proxy (xjcd)", [
        "GET /https://origin/path from a browser or curl",
        "XrdHttp + HTTP extension + PSS + plugin",
        "304, CGI, allow_origin, journals",
        "An empty allowlist denies every chained origin",
    ], "Proxy", PINK)

    s = new_slide(prs, 5, T)
    kicker(s, "Read path")
    title(s, "A hit requires the full request")
    lead(s, "Partial coverage is a miss. The origin is read and the bytes are appended. There is no block size and no prefetch.", y=1.9, h=0.7)
    card(s, 0.55, 2.75, 6.0, 3.8, "Hit - serve from journal", [
        "Interval tree: does [off, off+len) exist?",
        "pread fragments; optional crc32c",
        "No origin request",
        "HTTP: honour no-cache / validators",
    ])
    card(s, 6.8, 2.75, 6.0, 3.8, "Miss - fetch, then append", [
        "XrdCl File or HTTP GET",
        "pwrite journal; update filesize",
        "The next full-range read can hit",
        "PgRead always goes to the origin; pages are still journaled for later Read",
    ])

    s = new_slide(prs, 6, T)
    kicker(s, "I/O")
    title(s, "What is cached")
    table(s, 0.45, 2.0, 12.4, 4.7, [
        ["Operation", "Behaviour"],
        ["Read / ReadV", "Served only if the entire request is present; otherwise origin and append"],
        ["PgRead", "Always the origin. Pages are journaled for later Read"],
        ["Write / Truncate / PgWrite", "Passed through; the journal for that file is reset"],
        ["DirList / Stat", "Optional filesystem plugin (system = true)"],
        ["Identity", "On-disk path, not the client URL. The first process takes an exclusive flock"],
        ["Further processes, same file", "Read from the origin (no shared writer)"],
    ], col_w=[3.6, 8.8])

    s = new_slide(prs, 7, T)
    kicker(s, "On disk")
    title(s, "Append-only fragments, not blocks")
    card(s, 0.55, 2.05, 6.0, 3.5, "Layout", [
        "Default: <cache>/<host>:<port>/<path>/journal",
        "flat = true - SHA256 directory per file",
        "basepath = /store/ - omit the host; start at the federation prefix",
        "Listings .journalcache_list* · Stat .journalcache_stat",
        "$journal/.xjc/ is never evicted",
    ])
    card(s, 6.8, 2.05, 6.0, 3.5, "Journal file", [
        "jheader_t - magic, mtime, filesize, version",
        "header_t { offset, size } + data",
        "v2: uint32 crc32c trailer per fragment",
        "HTTP freshness in xattrs (etag, cache-control)",
        "Symlink journals are refused (O_NOFOLLOW)",
    ])
    add_tb(s, 0.55, 5.75, 12.2, 0.9, "Gaps are not filled. A later Read of a span that still has a gap misses until that span is complete.", size=15, color=MUTED)

    s = new_slide(prs, 8, T)
    kicker(s, "HTTP")
    title(s, "Freshness is stored with the journal")
    lead(s, "CGI and the HTTP extension pass origin validators into XrdCl. A matching If-None-Match or If-Modified-Since can return 304 before the file is opened.", y=1.9, h=0.8)
    card(s, 0.55, 2.85, 6.0, 3.7, "Cache-Control", [
        "no-store - do not journal",
        "no-cache - Stat this session before serving",
        "private - do not write the shared journal",
        "max-age / s-maxage / Expires - expire from cached-at",
    ])
    card(s, 6.8, 2.85, 6.0, 3.7, "HTTP extension", [
        "Maps request headers to CGI",
        "Optional HEAD to http_origin",
        "304 when a validator still matches",
        "Emits stored ETag / Last-Modified on continue",
        "Enforces allow_origin even in bypass",
    ])

    s = new_slide(prs, 9, T)
    kicker(s, "Forwarding")
    title(s, "The proxy stays closed until an origin is allowed")
    lead(s, "xjcd init writes configuration only. An empty allow_origin denies every chained URL (/https://…, /root://…, including N-hop unwrap).", y=1.9, h=0.8)
    card(s, 0.55, 2.85, 4.0, 3.7, "Unwrap and allowlist", [
        "Regex against the full URL, or exact hostname",
        "Invalid regex is rejected when added",
    ])
    card(s, 4.7, 2.85, 4.0, 3.7, "Deny - 403", [
        "Empty list, or no match",
        "xjcd init default",
    ], None, PINK)
    card(s, 8.85, 2.85, 4.0, 3.7, "Open and journal", [
        "xjc allow-origin add",
        "Enable pss.permit",
        "external_redirect → 302",
    ])

    s = new_slide(prs, 10, T)
    kicker(s, "Tools")
    title(s, "Three tools, one cache tree")
    card(s, 0.55, 2.1, 3.9, 4.4, "xjc", [
        "bypass, multi-origin, allowlist",
        "external redirects",
        "cleaner watermarks",
        "Reloads policy.conf on change",
    ], "Policy", GOLD)
    card(s, 4.7, 2.1, 3.9, 4.4, "xjcd", [
        "init - state, units, TLS ports",
        "render / validate / show",
        "No xjcd run; systemd starts xrootd",
        "--install-systemd installs the units",
    ], "Bootstrap", GOLD)
    card(s, 8.85, 2.1, 3.9, 4.4, "xjccleand", [
        "mtime eviction, skip .xjc/",
        "Disabled until xjc cleaner enable on",
        "Do not run together with in-plugin size=",
        "Removes journal, list, and stat files only",
    ], "Evict", PINK)

    s = new_slide(prs, 11, T)
    kicker(s, "XCache")
    title(s, "XCache is the XRootD site cache")
    lead(s, "XCache is XrdPfc (Proxy File Cache) in front of PSS. It is a shared server: files are stored in fixed-size blocks (default 256 KiB) and neighbouring blocks can be prefetched.", y=1.9, h=0.9)
    card(s, 0.55, 3.0, 4.0, 3.5, "Clients", ["xrdcp · ROOT · many jobs"])
    card(s, 4.7, 3.0, 4.0, 3.5, "XCache node", [
        "xrootd · pss · XrdPfc",
        "blocks + .cinfo · RAM + disk",
        "prefetch · purge · concurrent readers",
        "pfc.httpcc · ResourceMonitor",
    ])
    card(s, 8.85, 3.0, 4.0, 3.5, "Origin", ["WAN or disk server / federation"])

    s = new_slide(prs, 12, T)
    kicker(s, "Side by side")
    title(s, "Where data is stored")
    card(s, 0.55, 2.1, 6.0, 4.5, "JournalCache", [
        "In the client, or on a small xjcd proxy",
        "Exact ranges that were read",
        "Append-only journal · no prefetch",
        "One flock writer per file",
        "HTTP 304 · closed allowlist",
    ], "Client journal", GOLD)
    card(s, 6.8, 2.1, 6.0, 4.5, "XCache (XrdPfc)", [
        "Dedicated xrootd cache node",
        "Fixed blocks (256 KiB default)",
        "Prefetch adjacent blocks · .cinfo map",
        "Many clients, shared files",
        "Purge by age, usage, or plugin",
    ], "Site cache", PINK)

    s = new_slide(prs, 13, T)
    kicker(s, "1 : 1")
    title(s, "JournalCache and XCache")
    table(s, 0.35, 1.95, 12.6, 5.1, [
        ["", "JournalCache", "XCache (XrdPfc)"],
        ["Role", "XrdCl file plugin, optional HTTP extension", "Server OSS cache in front of PSS"],
        ["Sharing", "One process holds the journal lock", "Many clients on one cache node"],
        ["Unit of storage", "Exact read ranges (journal fragments)", "Fixed blocks (default 256 KiB)"],
        ["Prefetch", "None", "Adjacent blocks"],
        ["Read amplification", "Only the requested range", "Rounded up to a block, plus prefetch"],
        ["Writes", "Passed through; that journal is reset", "Optional write-through"],
        ["HTTP freshness", "CGI, extension 304, xattrs", "pfc.httpcc, Fsctl only-if-cached"],
        ["Deployment", "XRD_PLUGIN or xjcd init", "Dedicated xrootd cache node (pfc.* / pss.*)"],
    ], col_w=[2.6, 5.1, 4.9])

    s = new_slide(prs, 14, T)
    kicker(s, "Choice")
    title(s, "Recommended use")
    lead(s, "The two can run together. A site may operate XCache, while a client or xjcd instance still journals the ranges that client read.", y=1.9, h=0.7)
    card(s, 0.55, 2.75, 6.0, 3.85, "JournalCache", [
        "A laptop or worker that rereads the same sparse ranges",
        "ROOT or analysis that touches only selected branches or events",
        "A compact HTTP forwarder for /https://origin/path",
        "Workloads that must not fetch a whole file or a 256 KiB block",
        "A closed origin allowlist and a hot-reloadable policy (xjc)",
    ], "Use", GOLD)
    card(s, 6.8, 2.75, 6.0, 3.85, "XCache", [
        "A site or campus cache in front of a federation",
        "Many jobs reading the same files from one node",
        "Sequential or wide reads, where prefetch is useful",
        "An existing PSS deployment that needs a shared disk pool",
        "Production purge, gstream monitoring, and multi-reader files",
    ], "Use", PINK)

    s = new_slide(prs, 15, T)
    kicker(s, "Limits")
    title(s, "Current limits")
    table(s, 0.45, 2.0, 12.4, 4.85, [
        ["Constraint", "Implication"],
        ["Full-range hits only", "Any gap in the request is a miss, even if most bytes are local"],
        ["One flock per journal", "A second process on the same file does not share the writer"],
        ["No prefetch", "The next uncached range always requires an origin read"],
        ["Empty allowlist denies", "After xjcd init, add allow_origin before chained GETs succeed"],
        ["xjccleand starts disabled", "Enable it, or the tree grows until the disk is full"],
        ["listttl=0", "Listing and stat entries do not expire; only mutations invalidate them"],
        ["Not an XCache replacement", "The site-wide shared block cache remains XrdPfc"],
    ], col_w=[3.8, 8.6])

    s = new_slide(prs, 16, T)
    kicker(s, "Outlook")
    title(s, "Status")
    card(s, 0.55, 2.05, 3.9, 3.7, "Done", [
        "XrdCl file + optional FS plugin",
        "Range journal, crc32c, symlink rejection",
        "HTTP extension: 304, CGI, HEAD",
        "Closed allowlist and redirects",
        "xjc, xjcd, xjccleand",
    ], "Done", GOLD)
    card(s, 4.7, 2.05, 3.9, 3.7, "Not done", [
        "Partial-range hits",
        "A shared multi-process journal writer",
        "Prefetch (intentionally out of scope)",
        "A replacement for site-wide XCache",
        "An open-by-default forwarding proxy",
    ], "Not done", PINK)
    card(s, 8.85, 2.05, 3.9, 3.7, "Next", [
        "Propose the plugin for xrootd/master",
        "Keep XCache as the shared site cache",
        "Use JournalCache on clients and compact HTTP proxies",
        "Documentation: html/index.html, xjcd(1)",
    ], "Next", PINK)
    add_tb(s, 0.55, 5.95, 12.2, 0.7, "JournalCache stores the bytes a client has read, on that client. XCache remains the shared block cache in front of a federation.", size=15, color=MUTED)

    prs.save(OUT)
    print("Wrote", OUT)


if __name__ == "__main__":
    main()
